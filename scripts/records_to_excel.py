#!/usr/bin/env python3
"""Build the Excel working view from records/ (records -> Excel).

The inverse of scripts/excel_to_records.py. records/<ID>.yaml is the source of
truth; this script builds catalog/MFL_Dataset_Registry.xlsx **from scratch** so
the workbook stays a faithful import/working view (and a valid CI fallback).
Run it after any edit under records/:

    python3 scripts/records_to_excel.py

Why from scratch and not editing the existing workbook: the original workbook
was authored by Excel, and openpyxl cannot round-trip Excel-authored files
safely — on re-save Excel reported "We found a problem with some content" and
stripped features (2026-09-16). Known traps hit that day: (a) `cell.value =
None` does not clear a cell carrying a Hyperlink — the link text re-materializes
on save; (b) cross-sheet dropdowns stored as x14 extensions are silently
dropped; (c) a DataValidation formula written with a leading "=" marks the
whole package as corrupt. Building a clean openpyxl-native workbook avoids the
entire class of round-trip corruption.

- Sheet names and column order are exactly what mosaic_pipeline/registry.py
  reads (" Registry" with its leading space; 23 columns in COL order).
- The Instructions and Reference sheet content lives in
  spec/registry_workbook_content.json (extracted once from the original
  workbook, with the obsolete "ID auto-filled by formula" instruction
  corrected — positional ID formulas caused the 023/024 duplicates and the
  070-074 collisions).
- `living_landscape` bare codes are written back as the controlled
  "CODE — Name" values from the Reference sheet content (never hardcoded).
- Dropdowns are classic DataValidations via workbook defined names pointing at
  Reference ranges (computed from the JSON content, formulas without "=").
- MOSAIC styling: Ocean Teal header, White/Warm Sand banding, green/amber/red
  conditional formatting on Access level and Processing status.
- Verifies after saving: reloads the file, compares every Registry cell against
  the records, and parses every XML part of the package (exit 1 on failure).

The Excel path cannot carry the YAML-only forward fields (citation, doi, ...):
a --from-excel build is byte-identical in stac/ and differs in datasets.json
only where records carry a citation. That gap is by design (see SNAPSHOT_NOTE).
"""
from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path
from xml.dom import minidom

HERE = Path(__file__).resolve().parent.parent  # repo root (MOSAIC_catalog/)
sys.path.insert(0, str(HERE))

import openpyxl  # noqa: E402
import yaml  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

XLSX = HERE / "catalog" / "MFL_Dataset_Registry.xlsx"
RECORDS_DIR = HERE / "records"
CONTENT_JSON = HERE / "spec" / "registry_workbook_content.json"
SHEET = " Registry"  # leading space is intentional (matches registry.py)

# YAML key per Excel column, in column order 1..23 (matches registry.py COL).
ORDER = [
    "id", "title", "country", "theme", "living_landscape", "data_type",
    "spatial_resolution", "temporal_coverage", "source", "contact",
    "access_level", "license", "processing_status", "file_names",
    "current_location", "migration_status", "server_path", "description",
    "date_registered", "last_updated", "update_frequency", "download_url",
    "file_size",
]

# Reference-sheet list header -> (defined name, Registry column with dropdown)
LISTS = {
    "Countries": ("MOSAIC_Countries", 3),
    "MFL Themes": ("MOSAIC_Themes", 4),
    "Living Landscapes (canonical 2026-07-21)": ("MOSAIC_LLs", 5),
    "Data types": ("MOSAIC_DataTypes", 6),
    "CGIAR Centres / Sources": ("MOSAIC_Sources", 9),
    "Access levels": ("MOSAIC_Access", 11),
    "Migration status": ("MOSAIC_Migration", 16),
    "Update frequency": ("MOSAIC_UpdFreq", 21),
}

# MOSAIC palette
TEAL = "0D5C6B"
SAND = "F7F4EF"
NEARBLACK = "1C2B3A"

HEADER_FILL = PatternFill("solid", fgColor=TEAL)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=11, color=NEARBLACK)
GHOST_FONT = Font(name="Calibri", size=10, italic=True, color="8A8A8A")
BAND_FILL = PatternFill("solid", fgColor=SAND)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def load_records() -> list[dict]:
    return [yaml.safe_load(p.read_text().split("\n", 1)[1])
            for p in sorted(RECORDS_DIR.glob("*.yaml"))]


def build() -> tuple[list[dict], dict]:
    content = json.loads(CONTENT_JSON.read_text())
    recs = load_records()

    wb = openpyxl.Workbook()
    # --- Instructions ---
    ws = wb.active
    ws.title = "Instructions"
    for row in content["instructions"]["rows"]:
        ws.append(row)
    for col, w in content["instructions"]["widths"].items():
        ws.column_dimensions[col].width = w
    ws["A1"].font = Font(name="Georgia", size=16, bold=True, color=TEAL)
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value:
            ws.cell(r, 2).font = Font(name="Calibri", bold=True, color=TEAL)
            ws.cell(r, 2).alignment = TOP
        if ws.cell(r, 3).value:
            ws.cell(r, 3).font = BODY_FONT
            ws.cell(r, 3).alignment = WRAP

    # --- Reference ---
    ref = wb.create_sheet("Reference")
    ref_rows = content["reference"]["rows"]
    for row in ref_rows:
        ref.append(row)
    for col, w in content["reference"]["widths"].items():
        ref.column_dimensions[col].width = w
    for c in range(1, len(ref_rows[0]) + 1):
        if ref.cell(1, c).value:
            ref.cell(1, c).fill = HEADER_FILL
            ref.cell(1, c).font = HEADER_FONT
    for r in range(2, ref.max_row + 1):
        for c in range(1, len(ref_rows[0]) + 1):
            if ref.cell(r, c).value:
                ref.cell(r, c).font = BODY_FONT

    # canonical "CODE — Name" strings, straight from the Reference content
    ll_col = ref_rows[0].index("Living Landscapes (canonical 2026-07-21)")
    ll_disp = {}
    for row in ref_rows[1:]:
        v = row[ll_col] if ll_col < len(row) else None
        if v and " — " in str(v):
            ll_disp[str(v).split(" — ")[0].strip()] = str(v)

    # --- Registry ---
    reg = wb.create_sheet(SHEET)
    headers = content["registry"]["headers"]
    reg.append(headers)
    reg.append(content["registry"]["placeholder_row"])
    for col, w in content["registry"]["widths"].items():
        reg.column_dimensions[col].width = w
    for c in range(1, 24):
        cell = reg.cell(1, c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        reg.cell(2, c).font = GHOST_FONT
        reg.cell(2, c).alignment = WRAP
    reg.row_dimensions[1].height = 30

    for i, rec in enumerate(recs):
        r = 3 + i
        for c, key in enumerate(ORDER, start=1):
            v = rec.get(key)
            if key == "living_landscape":
                v = ll_disp.get(v, v)
            cell = reg.cell(r, c, v)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP if key == "description" else TOP
            if i % 2 == 1:
                cell.fill = BAND_FILL

    last = 2 + len(recs)
    reg.freeze_panes = "A3"
    reg.auto_filter.ref = f"A2:W{last}"

    # --- defined names + dropdowns (formula WITHOUT "=", see docstring) ---
    for header, (name, reg_col) in LISTS.items():
        c = ref_rows[0].index(header) + 1
        n_vals = sum(1 for row in ref_rows[1:]
                     if c - 1 < len(row) and row[c - 1] not in (None, ""))
        letter = get_column_letter(c)
        wb.defined_names.add(DefinedName(
            name, attr_text=f"Reference!${letter}$2:${letter}${1 + n_vals}"))
        dv = DataValidation(type="list", formula1=name, allow_blank=True)
        rl = get_column_letter(reg_col)
        dv.add(f"{rl}3:{rl}300")
        reg.add_data_validation(dv)

    # --- traffic-light conditional formatting (status fields) ---
    def light(col, value, fill, font):
        rl = get_column_letter(col)
        reg.conditional_formatting.add(
            f"{rl}3:{rl}300",
            CellIsRule(operator="equal", formula=[f'"{value}"'],
                       fill=PatternFill("solid", fgColor=fill),
                       font=Font(color=font)))
    for val, fl, fo in (("Open", "C6EFCE", "006100"),
                        ("Internal", "FFEB9C", "9C6500"),
                        ("Restricted", "FFC7CE", "9C0006")):
        light(11, val, fl, fo)
    for val, fl, fo in (("Validated", "C6EFCE", "006100"),
                        ("Processed", "FFEB9C", "9C6500"),
                        ("Raw", "FFC7CE", "9C0006")):
        light(13, val, fl, fo)

    wb.save(XLSX)
    return recs, ll_disp


def verify(recs: list[dict], ll_disp: dict) -> int:
    errs = 0
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[SHEET]
    for i, rec in enumerate(recs):
        r = 3 + i
        for c, key in enumerate(ORDER, start=1):
            want = rec.get(key)
            if key == "living_landscape":
                want = ll_disp.get(want, want)
            got = ws.cell(r, c).value
            if (want or None) != (got or None):
                errs += 1
                print(f"MISMATCH row {r} {key}: want={want!r} got={got!r}")
    for r in range(3 + len(recs), ws.max_row + 1):
        for c in range(1, 31):
            if ws.cell(r, c).value not in (None, ""):
                errs += 1
                print(f"GHOST row {r} col {c}: {ws.cell(r, c).value!r}")
    # package-level checks: every XML part parses; no "=" inside <formula1>
    z = zipfile.ZipFile(XLSX)
    for n in z.namelist():
        if n.endswith((".xml", ".rels")):
            try:
                minidom.parseString(z.read(n))
            except Exception as e:  # noqa: BLE001
                errs += 1
                print(f"BAD XML part {n}: {e}")
    if b"<formula1>=" in z.read(f"xl/worksheets/sheet3.xml"):
        errs += 1
        print('BAD: a <formula1> starts with "=" (Excel repair trigger)')
    return errs


def main() -> int:
    recs, ll_disp = build()
    errs = verify(recs, ll_disp)
    if errs:
        print(f"FAILED: {errs} problem(s) — do not commit this workbook.")
        return 1
    print(f"OK: {len(recs)} records written to {XLSX.name} and verified "
          f"(values + package XML).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
